# -*- coding: utf-8 -*-
"""填充 粮达e销-智能体搭建场景卡：保留模板结构，替换示例内容，保持原文档格式风格。"""
from copy import deepcopy
import docx
from docx.oxml.ns import qn

SRC = "/Users/tiger/Documents/汇报材料和ppt/贸易AI竞赛/粮达e销-智能体搭建场景卡-v3.docx"
DST = "/Users/tiger/Documents/汇报材料和ppt/贸易AI竞赛/粮达e销-智能体搭建场景卡-v4.docx"

doc = docx.Document(SRC)
paras = doc.paragraphs

# ---------- 工具函数 ----------

def strip_gray_overrides(pel):
    """去除 pPr>rPr 中示例文字的灰色/字号/加粗/斜体覆盖，并清理游离文本。"""
    pel.text = None
    for child in list(pel):
        child.tail = None
        if child.tag == qn('w:pPr'):
            rPr = child.find(qn('w:rPr'))
            if rPr is not None:
                for bad in ('w:color', 'w:sz', 'w:szCs', 'w:b', 'w:bCs', 'w:i', 'w:iCs'):
                    e = rPr.find(qn(bad))
                    if e is not None:
                        rPr.remove(e)
                if len(rPr) == 0:
                    child.remove(rPr)


def clone_clean(src_el):
    """深拷贝段落元素，去掉游离文本与示例 run，清理灰色覆盖。"""
    el = deepcopy(src_el)
    for child in list(el):
        if child.tag == qn('w:r') or child.tag == qn('w:hyperlink'):
            el.remove(child)
    strip_gray_overrides(el)
    return el


def add_run(el, text):
    r = el.makeelement(qn('w:r'), {})
    rPr = el.makeelement(qn('w:rPr'), {})
    fonts = el.makeelement(qn('w:rFonts'), {})
    fonts.set(qn('w:hint'), 'eastAsia')
    rPr.append(fonts)
    r.append(rPr)
    t = el.makeelement(qn('w:t'), {})
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    el.append(r)
    return el


def find_para(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise RuntimeError(f"未找到段落: {text}")


def paras_between(start_text, end_text):
    """返回两个标题之间的段落（不含标题本身）。"""
    ps = doc.paragraphs
    i = next(i for i, p in enumerate(ps) if p.text.strip() == start_text)
    j = next(i for i, p in enumerate(ps) if p.text.strip() == end_text)
    return ps[i + 1:j]


def delete_paras(ps):
    for p in ps:
        p._p.getparent().remove(p._p)


# ---------- 模板段落（在删除前先取好） ----------
tmpl_body = paras[25]._p         # Normal 正文（首行缩进、460行距）
tmpl_bullet_scene = paras[6]._p  # 场景章节项目符号（400行距）
tmpl_bullet_build = paras[27]._p  # 搭建思路项目符号（460行距）
tmpl_h2 = paras[26]._p           # Heading2（（一）整体架构）

anchor_value = find_para("三、智能体价值")
anchor_build = find_para("搭建思路")
anchor_steps = find_para("（二）搭建步骤")


def insert_before(anchor, blocks, bullet_tmpl):
    for kind, text in blocks:
        if kind == 'h2':
            el = clone_clean(tmpl_h2)
        elif kind == 'body':
            el = clone_clean(tmpl_body)
        elif kind == 'bullet':
            el = clone_clean(bullet_tmpl)
        else:
            raise ValueError(kind)
        add_run(el, text)
        anchor._p.addprevious(el)


# ================= 一、基本信息表格 =================
t0 = doc.tables[0]

def set_cell(cell, text):
    for p in list(cell.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    pel = cell.paragraphs[0]._p
    for child in list(pel):
        if child.tag == qn('w:r'):
            pel.remove(child)
    pel.text = None
    add_run(pel, text)

set_cell(t0.rows[0].cells[1], "李清（队长）、魏宇、罗正德、吴冠霖、周涛")
set_cell(t0.rows[1].cells[1], "粮达精灵团队")
set_cell(t0.rows[2].cells[1], "QoderWork")
set_cell(t0.rows[3].cells[1], "生产经营赛道")
set_cell(t0.rows[4].cells[1],
         "“粮达e销”——多智能体协同的粮食采购智能体：打造“粮掌柜 + 六位专业小二”数字员工团队，"
         "一句话下达采购目标，多小二按需协同形成可比较、可解释、可执行的采购方案，"
         "企业办事经验自动沉淀、跨任务复用。")

# ================= 二、智能体场景 =================
delete_paras(paras_between("二、智能体场景", "三、智能体价值"))

scene_blocks = [
    ('h2', "（一）场景背景与目标群体"),
    ('body', "粮达网（中粮招商局（深圳）粮食电子交易中心）是中粮集团与招商局集团联合打造的大宗农粮交易一站式服务平台，"
             "覆盖玉米、小麦、稻谷、大豆等品种的线上交易、电子签约、互联网金融、物流配送与行情资讯。"
             "粮达网已经解决了交易在线化与服务连接问题，但用户在发起或执行一笔粮贸业务时，"
             "仍需在行情、粮源、物流、综合成本、资金服务和风险之间人工寻找信息、反复核对并组织决策。"),
    ('body', "本场景首期聚焦饲料、养殖、粮食加工等下游用粮企业的原粮采购人员：他们有明确的业务目标——"
             "在对的时间、以合理的成本、安全地买到合适的粮，却往往不知道该进入平台的哪个功能，"
             "更难在多个约束条件下形成全局最优判断。"),
    ('h2', "（二）核心痛点"),
    ('body', "大宗粮食采购不是单点比价问题，行情、粮源、物流、资金、成本与履约风险相互影响，一线业务长期面临五大难题："),
    ('bullet', "信息散：行情、粮源、运力、资金与风险信息分散在多个系统、工作群和个人经验里，难以汇聚成统一的决策依据；"),
    ('bullet', "决策慢：一笔采购涉及多条专业线，靠人工多方沟通、反复核对，决策周期以天计，容易错过行情窗口；"),
    ('bullet', "比较难：“最低报价≠最低综合成本”，运费、损耗、质量折价与资金成本缺乏统一口径，多套方案难以横向比较；"),
    ('bullet', "经验断：采购经验依赖老师傅个人积累，难以沉淀为企业资产，人员变动直接带来决策水平波动；"),
    ('bullet', "风险后知后觉：供应方资质、合同条款与履约风险往往在出问题后才暴露，缺少前置的独立审核。"),
    ('h2', "（三）智能体总体设计"),
    ('body', "“粮达e销”把粮达网既有业务能力重新组织成 AI 原生的粮贸服务入口，按“用户不用先找功能、只需要叫小二；"
             "小二不用全部上场、只按需要入席”的理念，打造一支由 7 位 AI 数字员工组成的团队："
             "粮掌柜负责理解采购目标、拆解任务、推荐协作团队并汇总跨专业结论；瞻小二看行情、粮小二找粮源、"
             "运小二找物流、算小二算成本、钱小二找资金、安小二查风险。每位小二都能独立完成专业服务，"
             "也可围绕同一个采购目标按需协作。"),
    ('body', "系统坚持“AI 与确定性规则分工”：大模型负责自然语言理解、知识提炼与可读解释；硬条件过滤、成本归集、"
             "资金匹配、风险判定等关键计算由确定性规则完成，关键数字可核验；目标确认、团队确认、最终方案等关键节点"
             "设置人工决策闸门，AI 负责研究、组织、比较与提醒，最终业务决定始终由人做出。核心覆盖三类高频场景："),
    ('h2', "1. 多智能体协同采购决策"),
    ('bullet', "用户用一句话下达复杂目标（如“库存只够 5 天，未来 15 天要采购 200 吨二等玉米到潍坊，不能影响生产”），"
               "粮掌柜自动结构化目标、抽取品种、数量、交期与约束条件，缺少关键信息时主动向用户补齐；"),
    ('bullet', "粮掌柜推荐参与小二并逐一说明参与原因，用户确认后，瞻、粮、运、钱小二并行办理；"
               "算小二基于上游粮源与运输结果组合多套方案并测算到厂综合成本，安小二独立审核供应方履约证据与质量、交付风险；"),
    ('bullet', "自动检测跨专业冲突（如综合成本最低的方案到货时间不满足产线要求），并让相关小二做一次定向补充，"
               "最终形成推荐方案、备选方案、生效条件与行动清单，交由用户在决策闸门做最终确认。"),
    ('h2', "2. 单小二独立的专项服务"),
    ('bullet', "只看行情直接找瞻小二：输出当前判断、判断依据、对用户业务的影响与行动建议，并可设置价格与时间关注条件持续跟踪；"),
    ('bullet', "只找粮源、物流或资金直接找粮、运、钱小二：粮小二按品种、质量、交期等硬条件过滤，输出候选粮源对比与待核实事项；"
               "运小二输出路线方案的运费、时效与交付风险对比；钱小二按金额、期限与付款节点匹配资金产品，"
               "额度与期限判断全部由确定性规则完成，不做授信承诺；"),
    ('bullet', "只算成本或查风险直接找算、安小二：算小二把不同方案转换到同一到厂成本口径，并标明确认项与估算项；"
               "安小二输出风险分级、判断依据与核验动作；任何单项服务都可按需升级为多小二协作任务。"),
    ('h2', "3. 企业知识沉淀与复用"),
    ('bullet', "从用户确认的办事结果中自动提炼企业事实、业务偏好、决策经验与风险规则，形成随业务自动生长的“企业知识大脑”；"),
    ('bullet', "小二办理新任务时主动召回并应用相关经验，页面上明确展示引用了哪条经验、为什么适用、对方案排序产生了什么影响；"),
    ('bullet', "用户可拒绝本次采用、编辑适用范围或停用某条知识，引用轨迹全程可查、可纠偏，让企业记忆真正“从业务中来、到业务中去”。"),
]
insert_before(anchor_value, scene_blocks, tmpl_bullet_scene)

# ================= 三、智能体价值 =================
delete_paras(paras_between("三、智能体价值", "搭建思路"))

value_blocks = [
    ('body', "“粮达e销”直击大宗粮食采购“信息散、决策慢、比较难、经验断、风险后知后觉”的痛点，"
             "价值体现在提效、降本、风控、助决策、沉淀经验五个方面："),
    ('bullet', "提效：采购方案从“多天多轮沟通”变为“一句话下达、分钟级产出”。行情研判、粮源寻源、物流规划、资金匹配"
               "由多小二并行完成，成本比较与风险审核一次成稿，把采购员从重复的信息搜集与核对中解放出来，"
               "聚焦议价与供应链关系经营。"),
    ('bullet', "降本：算小二以统一的综合成本口径，让运费、损耗、质量折价、资金成本等隐性成本可见、可比，"
               "避免“最低报价”误导采购决策；物流路线与发运安排优化进一步降低交付成本与资金占用。"),
    ('bullet', "风控：安小二对供应方、合同条款与履约方案独立审核、不为迎合推荐方案而弱化风险；"
               "跨专业冲突自动检测并显性呈现，让风险在确认前被拦截，显著降低人工疏漏与合规风险。"),
    ('bullet', "助决策：每个结论都带判断依据，推荐方案、备选方案与生效条件并列呈现，“为什么选这个方案”一目了然，"
               "让采购决策更快、更有依据、可追溯。"),
    ('bullet', "沉淀经验：每一次确认过的办事结果都自动沉淀为企业知识并在后续任务中复用，采购经验第一次从"
               "“老师傅的脑子”变成“企业的资产”，新人可以快速复用团队积累的决策水平。"),
    ('body', "从可量化角度看，采购方案形成周期可由天级压缩到小时级乃至分钟级，多方案比较覆盖率与风险前置发现率显著提升，"
             "企业知识资产的增长可长期度量；“多智能体协同 + 人工决策闸门 + 企业记忆复用”的模式不绑定具体品种与区域，"
             "可横向复制到不同品类与业务环节，规模化收益明显。"),
]
insert_before(anchor_build, value_blocks, tmpl_bullet_scene)

# ================= 搭建思路 =================
# 开头总述段（替换原示例段）
old_intro = find_para("（示例）依托 QoderWork 的对话编排与技能调度能力，采用“数据接入 + 能力编排 + 模板沉淀”的方式分阶段落地，先跑通单一场景再横向扩展。")
new_intro = clone_clean(tmpl_body)
add_run(new_intro, "“粮达e销”依托 QoderWork 的 AI 编程与智能体搭建能力，按“AI 理解 + 规则计算 + 人做决定 + 经验沉淀”"
                   "的总体思路分阶段落地：先跑通“多小二协同采购”主闭环，再横向扩展更多品种与场景。")
old_intro._p.addprevious(new_intro)
old_intro._p.getparent().remove(old_intro._p)

# 架构层（替换原 4 条示例）
delete_paras(paras_between("（一）整体架构", "（二）搭建步骤"))

arch_blocks = [
    ('bullet', "交互层：科技风 React 工作台，以粮掌柜数字协作场为统一入口，用户通过自然语言对话下达任务，"
               "实时查看多小二办理过程，并在决策闸门完成关键确认。"),
    ('bullet', "编排层：基于 LangChain / LangGraph 构建多智能体编排图，调度“目标解析 → 团队推荐 → 并行办理 → 依赖办理 "
               "→ 冲突检测 → 综合建议”全流程，支持任务挂起、恢复与过程事件回放。"),
    ('bullet', "能力层：七位数字员工各自承载行情研判、粮源寻采、物流规划、成本测算、资金匹配、风险审核等专业技能，"
               "小二之间通过结构化任务交接完成独立服务与协作升级。"),
    ('bullet', "知识层：从确认过的办事记录中提炼企业事实、业务偏好、决策经验与风险规则，经 Mem0 + Milvus 语义召回、"
               "MySQL 检索兜底，知识引用全程可追踪、可纠偏。"),
    ('bullet', "数据与规则层：FastAPI + MySQL 统一管理行情、粮源、物流线路、金融产品与办事记录；硬条件过滤、成本归集"
               "与风险判定由确定性规则完成，确保关键计算不由模型生成、结果可核验。"),
]
insert_before(anchor_steps, arch_blocks, tmpl_bullet_build)

# ================= 搭建步骤表格 =================
t1 = doc.tables[1]
steps = [
    ("步骤一 场景拆解与角色定义",
     "围绕大宗粮食采购业务，把采购决策拆解为行情、粮源、物流、成本、资金、风险六大专业域；"
     "定义七位数字员工的职责与产出标准，确立“能独立完成不强行协作、按需邀请、冲突显性化、关键节点人工把关”的协作规则。"),
    ("步骤二 数据底座与确定性规则",
     "构建行情、粮源、物流线路、资金产品等业务数据，落地硬条件过滤、综合成本归集、资金产品匹配、风险分级等确定性计算，"
     "确保价格、成本、额度等关键数字全部由规则产出，模型只做理解与解释。"),
    ("步骤三 多智能体编排开发",
     "基于 LangGraph 实现粮掌柜目标结构化与团队推荐，瞻、粮、运、钱小二并行办理，算小二依赖测算，安小二独立审核，"
     "以及跨专业冲突检测与定向补充，形成“一个目标进、一套可执行方案出”的主链路。"),
    ("步骤四 人工闸门与企业知识闭环",
     "在目标确认、团队确认、最终方案三处设置人工决策闸门，确保 AI 不越权；从确认过的办事记录中自动提炼企业知识，"
     "经 Mem0 + Milvus 语义召回，小二在新任务中主动引用经验，引用原因与轨迹可见，支持一键拒绝与停用。"),
    ("步骤五 场景搭建与端到端验证",
     "开发科技风多智能体协作工作台，实现办理过程可视化、冲突呈现、知识引用交互与决策闸门确认；"
     "完成“目标输入 → 协同办理 → 方案确认 → 知识沉淀 → 经验复用”全链路验证，形成可复现、可演示的业务闭环。"),
]

for i, (step, desc) in enumerate(steps, start=1):
    set_cell(t1.rows[i].cells[0], step)
    set_cell(t1.rows[i].cells[1], desc)

# ================= 表格后补充“当前进展” =================
progress = clone_clean(tmpl_body)
add_run(progress, "当前进展：项目已完成端到端可运行的原型系统，覆盖粮掌柜目标解析与团队推荐、六位专业小二办理、"
                  "冲突检测与综合建议、人工决策闸门与企业知识大脑等全部核心模块，"
                  "“多小二协同采购”与“企业知识闭环”两条演示路线均可完整复现，代码仓库："
                  "https://github.com/growth-ux/liandgawang-plus 。")
doc.tables[1]._tbl.addnext(progress)

doc.save(DST)
print("已保存:", DST)
